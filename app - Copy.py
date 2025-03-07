import os
from flask import Flask, request, render_template, send_file
import pdfplumber
import docx
import re
import openpyxl

app = Flask(__name__)

# Set up the path for storing uploaded files (using Render's persistent disk)
UPLOAD_FOLDER = '/mnt/data/uploads'  # Render's persistent storage path
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

# Function to extract email
def extract_email(text):
    email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
    email = re.findall(email_pattern, text)
    return email[0] if email else None

# Function to extract phone number
def extract_phone(text):
    phone_pattern = r'\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}|\d{10,15}'
    phone = re.findall(phone_pattern, text)
    return phone[0] if phone else None

# Function to extract name (assuming the name is in the first line/paragraph)
def extract_name(text):
    lines = text.splitlines()
    for line in lines:
        if line.strip():
            return line.strip()
    return None

# Function to read PDF file
def read_pdf(file_path):
    with pdfplumber.open(file_path) as pdf:
        text = ''
        for page in pdf.pages:
            text += page.extract_text()
    return text

# Function to read DOCX file
def read_docx(file_path):
    doc = docx.Document(file_path)
    text = ''
    for para in doc.paragraphs:
        text += para.text
    return text

# Function to process resumes and extract data
def process_resumes(folder_path):
    resume_data = []
    for file_name in os.listdir(folder_path):
        if file_name.endswith('.pdf'):
            text = read_pdf(os.path.join(folder_path, file_name))
        elif file_name.endswith('.docx'):
            text = read_docx(os.path.join(folder_path, file_name))
        else:
            print(f"Unsupported file type: {file_name}")
            continue

        # Extract name, email, and phone number from the text
        name = extract_name(text)
        email = extract_email(text)
        phone = extract_phone(text)
        
        # Store the extracted data
        resume_data.append([name, email, phone])

    return resume_data

# Function to create and save data into an Excel file
def create_excel(data, output_file):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resume Data"

    headers = ["Name", "Email", "Phone Number"]
    ws.append(headers)
    for row in data:
        ws.append(row)

    wb.save(output_file)

@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        # Get the uploaded file
        uploaded_files = request.files.getlist('file')

        # Create a temporary directory inside the uploads folder to store the files
        temp_folder = os.path.join(app.config['UPLOAD_FOLDER'], 'temp')

        # Ensure the directory exists
        if not os.path.exists(temp_folder):
            try:
                os.makedirs(temp_folder)
            except PermissionError as e:
                return f"PermissionError: {e} - Please check your permissions."

        # Process resumes and extract data
        resume_data = []
        
        for file in uploaded_files:
            # Save the uploaded files in the temporary directory
            file_path = os.path.join(temp_folder, file.filename)
            file.save(file_path)
            
            # Process the files to extract data
            resume_data.extend(process_resumes(temp_folder))

        # Create the Excel file
        output_file = os.path.join(app.config['UPLOAD_FOLDER'], 'resume_data.xlsx')
        create_excel(resume_data, output_file)

        # Clean up the temporary folder
        for file in os.listdir(temp_folder):
            os.remove(os.path.join(temp_folder, file))
        os.rmdir(temp_folder)

        # Send the generated Excel file to the user
        return send_file(output_file, as_attachment=True)

    return render_template('index.html')

if __name__ == "__main__":
    # Bind to the port specified by Render
    port = os.getenv("PORT", 5000)  # Use Render's port or fallback to 5000
    app.run(host="0.0.0.0", port=int(port), debug=True)
