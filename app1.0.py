import os
import tempfile
from flask import Flask, request, render_template, send_file
import pdfplumber
import docx
import re
import openpyxl
from datetime import datetime



# Mapping country names to nationalities
country_to_nationality = {
    "United States": "American",
    "USA": "American",
    "India": "Indian",
    "Canada": "Canadian",
    "United Kingdom": "British",
    "UK": "British",
    "Australia": "Australian",
    "Germany": "German",
    "France": "French",
    "Spain": "Spanish",
    "Italy": "Italian",
    "China": "Chinese",
    "Japan": "Japanese",
    "South Korea": "Korean",
    "Brazil": "Brazilian",
    "Mexico": "Mexican",
    "Russia": "Russian",
    "Netherlands": "Dutch",
    "Turkey": "Turkish",
    "Sweden": "Swedish",
    "Norway": "Norwegian",
    "Denmark": "Danish",
    "Finland": "Finnish",
    "Switzerland": "Swiss",
    "South Africa": "South African",
    "Argentina": "Argentinian",
    "Egypt": "Egyptian",
    "Saudi Arabia": "Saudi",
    "United Arab Emirates": "Emirati",
    "UAE": "Emirati",
}


app = Flask(__name__)

# Use tempfile to handle temporary directory creation
UPLOAD_FOLDER = tempfile.mkdtemp()  # Creates a temporary directory in a safe location
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

# Function to extract email
def extract_email(text):
    email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
    email = re.findall(email_pattern, text)
    return email[0] if email else None

# Function to extract phone number
def extract_phone(text):
    phone_pattern = r'\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}|\d{10,15}'
    phone = re.findall(phone_pattern, text)
    return phone[0] if phone else None

# Function to extract name (assuming the name is in the first line/paragraph)
def extract_name(text):
    lines = text.splitlines()
    for line in lines:
        if line.strip():
            return line.strip()
    return None
    
    
 # Function to extract nationality based on country name
def extract_nationality(text):
    for country, nationality in country_to_nationality.items():
        if re.search(rf'\b{country}\b', text, re.IGNORECASE):  # Ensures whole-word match
            return nationality
    return "Not Found"   

# Function to read PDF file
def read_pdf(file_path):
    with pdfplumber.open(file_path) as pdf:
        text = ''
        for page in pdf.pages:
            text += page.extract_text()
    return text

# Function to read DOCX file
def read_docx(file_path):
    doc = docx.Document(file_path)
    text = ''
    for para in doc.paragraphs:
        text += para.text
    return text

# Function to process a single resume and extract data
def process_resume(file_path):
    # Read the file and extract text
    text = ''
    if file_path.endswith('.pdf'):
        text = read_pdf(file_path)
    elif file_path.endswith('.docx'):
        text = read_docx(file_path)
    else:
        print(f"Unsupported file type: {file_path}")
        return []

    # Extract name, email, and phone number from the text
    name = extract_name(text)
    email = extract_email(text)
    phone = extract_phone(text)
    nationality = extract_nationality(text)
    
    # Return the extracted data
    return [(name, email, phone, nationality)]

# Function to create and save data into an Excel file
def create_excel(data, output_file):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resume Data"

    headers = ["Name", "Email", "Phone Number", "Nationality"]
    ws.append(headers)
    for row in data:
        ws.append(row)

    wb.save(output_file)

@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        # Get the uploaded file(s)
        uploaded_files = request.files.getlist('file')

        # Initialize list to hold all extracted resume data
        all_resume_data = []
        
        # Process each uploaded file individually
        for file in uploaded_files:
            # Save the uploaded file in the temporary upload folder
            file_path = os.path.join(app.config['UPLOAD_FOLDER'], file.filename)
            file.save(file_path)
            
            # Process the file to extract data and add it to the overall data
            resume_data = process_resume(file_path)
            all_resume_data.extend(resume_data)

        # Get current datetime for dynamic file naming
        current_datetime = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        
        # Create dynamic output file name with the current datetime
        output_file = os.path.join(app.config['UPLOAD_FOLDER'], f'resumedata_{current_datetime}.xlsx')

        # Create the Excel file with extracted data
        create_excel(all_resume_data, output_file)

        # Send the generated Excel file to the user
        return send_file(output_file, as_attachment=True)

    return render_template('index.html')

if __name__ == "__main__":
    # Bind to the port specified by Render
    port = os.getenv("PORT", 5000)  # Use Render's port or fallback to 5000
    app.run(host="0.0.0.0", port=int(port), debug=True)
