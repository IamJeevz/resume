import pdfplumber
import docx
import re
import openpyxl
import os


# Function to extract email
def extract_email(text):
    email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
    email = re.findall(email_pattern, text)
    return email[0] if email else None


# Function to read PDF file
def read_pdf(file_path):
    with pdfplumber.open(file_path) as pdf:
        text = ''
        for page in pdf.pages:
            text += page.extract_text()  # Fix the typo here: 'text =+ page.extract_text()' to 'text += page.extract_text()'
    return text


# Function to read DOCX file
def read_docx(file_path):
    doc = docx.Document(file_path)
    text = ''
    for para in doc.paragraphs:
        text += para.text
    return text


# Function to process resumes and extract data
def process_resumes(folder_path):
    resume_data = []
    for file_name in os.listdir(folder_path):
        if file_name.endswith('.pdf'):
            text = read_pdf(os.path.join(folder_path, file_name))
        elif file_name.endswith('.docx'):
            text = read_docx(os.path.join(folder_path, file_name))
        else:
            print(f"Unsupported file type: {file_name}")
            continue

        # Extract email from the text
        email = extract_email(text)
        resume_data.append([email])

    return resume_data


# Function to create and save data into an Excel file
def create_excel(data, output_file):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resume Data"

    headers = ["Email"]
    ws.append(headers)
    for row in data:
        ws.append(row)

    wb.save(output_file)


# Main function
def main():
    folder_path = "F:/Python Projects/Scanner"  # Make sure the path is correct
    output_file = "text.xlsx"

    resume_data = process_resumes(folder_path)
    create_excel(resume_data, output_file)

    print(f"Saved successfully to {output_file}")


if __name__ == "__main__":
    main()
