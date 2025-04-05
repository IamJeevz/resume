import os
import tempfile
from flask import Flask, request, render_template, send_file
import pdfplumber
import docx
import re
import openpyxl
from datetime import datetime
from difflib import SequenceMatcher

# Mapping country names to nationalities
country_to_nationality = {
    "United States": "American", "USA": "American", "India": "Indian", "Canada": "Canadian",
    "United Kingdom": "British", "UK": "British", "Australia": "Australian", "Germany": "German",
    "France": "French", "Spain": "Spanish", "Italy": "Italian", "China": "Chinese", "Japan": "Japanese",
    "South Korea": "Korean", "Brazil": "Brazilian", "Mexico": "Mexican", "Russia": "Russian",
    "Netherlands": "Dutch", "Turkey": "Turkish", "Sweden": "Swedish", "Norway": "Norwegian",
    "Denmark": "Danish", "Finland": "Finnish", "Switzerland": "Swiss", "South Africa": "South African",
    "Argentina": "Argentinian", "Egypt": "Egyptian", "Saudi Arabia": "Saudi", 
    "United Arab Emirates": "Emirati", "UAE": "Emirati"
}

app = Flask(__name__)

# Use a temporary directory for file uploads
UPLOAD_FOLDER = tempfile.mkdtemp()
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

# Function to check name similarity
def name_similarity(extracted_name, file_name):
    ratio = SequenceMatcher(None, extracted_name.lower(), file_name.lower()).ratio()
    print(f"Similarity Ratio: {ratio:.2f}")
    return ratio

# Function to extract email
def extract_email(text):
    email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
    email = re.findall(email_pattern, text)
    return email[0] if email else None

def extract_phone(text):
    phone_pattern = r'\+?\d{1,3}[-.\s]?\d{1,4}[-.\s]?\d{2,4}[-.\s]?\d{2,4}[-.\s]?\d{2,4}'
    phone = re.findall(phone_pattern, text)
    return phone[0] if phone else None


# Function to extract name (assuming it is in the first non-empty line)
def extract_name(text):
    lines = text.splitlines()
    for line in lines:
        if line.strip():
            return line.strip()
    return None

# Function to extract nationality based on country mention
def extract_nationality(text):
    found_countries = []
    for country, nationality in country_to_nationality.items():
        if re.search(rf'\b{country}\b', text, re.IGNORECASE):
            found_countries.append(nationality)

    nationality_match = re.search(r'Nationality[:\-]?\s*(\w+)', text, re.IGNORECASE)
    if nationality_match:
        return nationality_match.group(1).capitalize()

    if len(found_countries) > 1:
        return ", ".join(set(found_countries))

    return found_countries[0] if found_countries else "Not Found"

# Function to read PDF file
def read_pdf(file_path):
    with pdfplumber.open(file_path) as pdf:
        text = ''.join(page.extract_text() for page in pdf.pages if page.extract_text())
    return text

# Function to read DOCX file
def read_docx(file_path):
    doc = docx.Document(file_path)
    return '\n'.join(para.text for para in doc.paragraphs)

# Function to process a single resume and extract data
def process_resume(file_path):
    file_name = os.path.splitext(os.path.basename(file_path))[0]  # Extract file name without extension

    # Read the file and extract text
    text = ''
    if file_path.endswith('.pdf'):
        text = read_pdf(file_path)
    elif file_path.endswith('.docx'):
        text = read_docx(file_path)
    else:
        print(f"Unsupported file type: {file_path}")
        return []

    # Extract details
    name = extract_name(text)
    email = extract_email(text)
    phone = extract_phone(text)
    nationality = extract_nationality(text)

    # Check name similarity with file name
    name_match = "Match" if name and name_similarity(name, file_name) > 0.7 else "Mismatch"

    # Return extracted data
    return [(name, email, phone, nationality)]

# Function to create and save data into an Excel file
def create_excel(data, output_file):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resume Data"

    headers = ["Name", "Email", "Phone Number", "Nationality"]
    ws.append(headers)
    for row in data:
        ws.append(row)

    wb.save(output_file)

@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        uploaded_files = request.files.getlist('file')
        all_resume_data = []

        for file in uploaded_files:
            file_path = os.path.join(app.config['UPLOAD_FOLDER'], file.filename)
            file.save(file_path)
            
            # Process file and extract data
            resume_data = process_resume(file_path)
            all_resume_data.extend(resume_data)

        # Generate unique output filename
        current_datetime = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        output_file = os.path.join(app.config['UPLOAD_FOLDER'], f'resumedata_{current_datetime}.xlsx')

        # Save extracted data into Excel
        create_excel(all_resume_data, output_file)

        return send_file(output_file, as_attachment=True)

    return render_template('index.html')

if __name__ == "__main__":
    port = os.getenv("PORT", 5000)  # Use Render's port or default to 5000
    app.run(host="0.0.0.0", port=int(port), debug=True)
